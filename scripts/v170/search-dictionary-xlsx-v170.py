"""V170 search dictionary workbook: export for review and editing, import back.

The synonym dictionary is kept in the repository as JSON
(src/data/search/searchSynonymsV170.json) because the site reads it. People
edit it in a workbook; this script moves it between the two.

  python3 scripts/v170/search-dictionary-xlsx-v170.py export OUT.xlsx
      Sheets: 안내, 동의어 사전 (editable), 주요 현황 규칙 (reference,
      with the values each country shows now).
  python3 scripts/v170/search-dictionary-xlsx-v170.py import IN.xlsx [--check]
      Reads 동의어 사전 back into the JSON (schemaVersion and note kept).
      Refuses a term listed in two groups, a duplicate group id, or an empty
      group. --check only validates and prints what would change.

After an import: node scripts/v170/build-search-v170.mjs (recounts the panel),
then the unit tests (src/data/search/searchMatchV170.test.ts).
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SYNONYMS = ROOT / "src/data/search/searchSynonymsV170.json"
RULES = ROOT / "scripts/v170/topic-rules-v170.json"
TOPICS = {iso3: ROOT / f"public/data/search/v170/topics-{iso3}.json" for iso3 in ("VNM", "BGD")}
CARDS = ROOT / "src/data/spec/datasetCardSpecV159.json"
TECH_CATALOG = ROOT / "src/data/climateTechnologyCatalog.ts"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F3A5F")
EDIT_FILL = PatternFill("solid", start_color="FFF7D6")
THIN = Side(style="thin", color="C8CFD6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SYN_SHEET = "동의어 사전"
SYN_HEADERS = ["그룹 ID", "대표어", "검색어 (한글)", "검색어 (영문·약어)", "기후기술 번호", "기후기술명", "비고"]
ASCII = re.compile(r"^[\x20-\x7e]+$")


def key(term: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", term).lower())


def tech_names() -> dict[str, str]:
    text = TECH_CATALOG.read_text(encoding="utf-8")
    names = re.findall(r'nameKo: "([^"]+)"', text)
    return {f"{i:02d}": name for i, name in enumerate(names, start=1)}


def split_terms(cell) -> list[str]:
    if cell is None:
        return []
    return [part.strip() for part in re.split(r"[,\n]", str(cell)) if part.strip()]


# ------------------------------------------------------------------ layout
def write_table(ws, headers, rows, widths, editable_cols=()):
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if c in editable_cols:
                cell.fill = EDIT_FILL
    for c, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def guide_sheet(wb):
    ws = wb.create_sheet("안내")
    lines = [
        ("데이터 찾기 검색사전", True),
        ("", False),
        ("1. 동의어 사전(수정 가능)", True),
        ("○ 같은 그룹의 단어는 같은 뜻으로 보고 함께 검색함. 어느 단어로 검색해도 그룹 전체를 검색함(양방향)", False),
        ("○ 등록 대상: 같은 뜻의 다른 표현, 표기 변형, 영문 명칭·약어(예: 태양광 = 태양광발전 = PV = photovoltaic)", False),
        ("○ 등록 제외: 상위·하위 개념(재생에너지 ⊃ 태양광), 관련어(태양광 ≠ 일사량). 등록 시 관계없는 데이터가 '데이터명·설명 일치'로 분류됨", False),
        ("○ 띄어쓰기·대소문자는 자동으로 무시하므로 별도로 등록하지 않음(태양광 발전 = 태양광발전, PV = pv)", False),
        ("○ 영문·약어는 단어 단위로만 검색함(PV로 'PVC'를 찾지 않음). 2글자 약어는 다른 뜻과 겹치지 않는지 확인 후 등록함", False),
        ("○ 한 단어는 한 그룹에만 등록함. 같은 단어가 두 그룹에 있으면 반영(import) 단계에서 오류로 처리됨", False),
        ("○ 노란 칸(대표어·검색어·비고)만 수정함. 검색어는 쉼표로 구분하며, 대표어는 그룹 이름이므로 검색어 칸에도 기재해야 검색됨", False),
        ("○ 새 그룹은 맨 아래 행에 그룹 ID(영문 소문자·숫자·하이픈)와 함께 추가하며, 검색어는 2개 이상 기재함", False),
        ("○ 그룹 ID는 '주요 현황 규칙'에서 참조하므로 기존 ID는 변경하지 않음", False),
        ("○ 기후기술 번호·기후기술명: 플랫폼 38대 기후기술 분류 중 같은 기술을 가리키는 그룹에만 기재함(참고용이며 검색에는 사용하지 않음)", False),
        ("○ 번호가 빈 14개 그룹(홍수·가뭄·탄소시장·ODA·핵심광물 등)은 기후기술이 아닌 기후위험·탄소시장·정책·재원 주제이므로 비워 둠. 하나의 위험이 여러 기술과 관련되므로 특정 번호로 지정하지 않음", False),
        ("", False),
        ("2. 주요 현황 규칙(참고용)", True),
        ("○ 검색어가 주제에 해당하면 검색 결과 위에 주제별 주요 현황을 표시함(국가 1개 선택, 그 밖의 필터 미선택 시)", False),
        ("○ 표시값은 각 데이터의 원자료를 기준으로 산출한 값임. 값이 없는 행은 표시하지 않으며, 3행 미만인 주제는 주요 현황을 표시하지 않음", False),
        ("○ 산출 방식: 대표값(조건 문구가 있는 경우만) / 지표 최신값(지표 1개, 합산 안 함) / 시설 집계(시설 수·용량 합계) / 수록 내용 건수(주제어가 들어간 건수)", False),
        ("○ 규칙 변경(주제·항목 추가)은 개발 담당자가 scripts/v170/topic-rules-v170.json에 반영함", False),
        ("", False),
        ("3. 반영 절차(개발 담당)", True),
        ("① python3 scripts/v170/search-dictionary-xlsx-v170.py import <이 파일> --check  (검증)", False),
        ("② python3 scripts/v170/search-dictionary-xlsx-v170.py import <이 파일>  (동의어 사전 반영)", False),
        ("③ node scripts/v170/build-search-v170.mjs  (검색용 데이터 파일·주요 현황 수치 재생성)", False),
        ("④ npm run test:unit  (단위 시험) → PR", False),
    ]
    for text, bold in lines:
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, bold=bold, size=12 if bold and ws.max_row == 1 else 11)
    ws.column_dimensions["A"].width = 120
    return ws


def synonym_rows(groups, techs):
    rows = []
    for group in groups:
        ko = [t for t in group["terms"] if not ASCII.match(t)]
        en = [t for t in group["terms"] if ASCII.match(t)]
        tech = group.get("techId") or ""
        rows.append([group["id"], group["label"], ", ".join(ko), ", ".join(en), tech, techs.get(tech, ""), group.get("note", "")])
    return rows


def rule_text(rule) -> tuple[str, str]:
    kind = rule.get("type")
    if kind == "card":
        words = "' 또는 '".join(rule["require"].split("|")) if rule.get("require") else ""
        cond = f"대표값에 '{words}' 포함 시" if words else "-"
        return "대표값", cond
    if kind == "indicatorLatest":
        return "지표 최신값(지표 1개, 합산 안 함)", f"지표: {rule.get('label', '')}"
    if kind == "entityAgg":
        return "시설 집계(시설 수·용량 합계)", "주연료가 주제어(동의어 포함)인 시설"
    if kind == "recordMatch":
        return "수록 내용 건수", "주제어(동의어 포함)가 들어간 수록 내용"
    return kind or "", ""


def rules_rows(rules, cards, shown):
    rows = []
    for topic in rules["topics"]:
        for order, rule in enumerate(topic["rows"], start=1):
            options = rule.get("oneOf") or [rule]
            for index, option in enumerate(options):
                method, cond = rule_text(option)
                if len(options) > 1:
                    method = f"({index + 1}순위) {method}"
                eid = option["elementId"]
                values = []
                for iso3 in ("VNM", "BGD"):
                    row = shown[iso3].get((topic["id"], eid))
                    values.append(f"{row['value']} ({row['sub']})" if row and (len(options) == 1 or row.get("basis") == option["type"]) else "")
                rows.append([
                    topic["label"], ", ".join(topic["groups"]), order, rule["role"], eid,
                    cards.get(eid, ""), method, cond, *values,
                ])
    return rows


def export(out: Path):
    syn = json.loads(SYNONYMS.read_text(encoding="utf-8"))
    rules = json.loads(RULES.read_text(encoding="utf-8"))
    cards = {row["elementId"]: row.get("baseName", "") for row in json.loads(CARDS.read_text(encoding="utf-8"))["rows"]}
    shown = {}
    for iso3, path in TOPICS.items():
        data = json.loads(path.read_text(encoding="utf-8"))
        shown[iso3] = {(t["id"], r["elementId"]): r for t in data["topics"] for r in t["rows"]}
    techs = tech_names()

    wb = Workbook()
    wb.remove(wb.active)
    guide_sheet(wb)
    ws = wb.create_sheet(SYN_SHEET)
    write_table(ws, SYN_HEADERS, synonym_rows(syn["groups"], techs), [18, 14, 42, 42, 12, 26, 30], editable_cols=(2, 3, 4, 7))
    ws = wb.create_sheet("주요 현황 규칙")
    write_table(
        ws,
        ["주제", "동의어 그룹", "순서", "항목", "데이터 ID", "데이터명", "산출 방식", "조건", "베트남 표시값", "방글라데시 표시값"],
        rules_rows(rules, cards, shown),
        [12, 22, 7, 14, 10, 30, 30, 30, 40, 40],
    )
    wb.save(out)
    print(f"wrote {out} ({len(syn['groups'])} groups, {len(rules['topics'])} topics)")


def dump_synonyms(syn, groups) -> str:
    """The repository layout: one group per line."""
    def one(group):
        inner = json.dumps(group, ensure_ascii=False, separators=(", ", ": "))[1:-1]
        return "    { " + inner + " }"
    head = json.dumps({"schemaVersion": syn["schemaVersion"], "note": syn["note"]}, ensure_ascii=False, indent=2)[:-2]
    return head + ',\n  "groups": [\n' + ",\n".join(one(g) for g in groups) + "\n  ]\n}\n"


def import_(src: Path, check: bool):
    syn = json.loads(SYNONYMS.read_text(encoding="utf-8"))
    ws = load_workbook(src, data_only=True)[SYN_SHEET]
    header = [str(c.value or "").strip() for c in ws[1]]
    if header[: len(SYN_HEADERS)] != SYN_HEADERS:
        sys.exit(f"'{SYN_SHEET}' header changed: {header}")
    previous = {g["id"]: g for g in syn["groups"]}
    groups, owner, errors = [], {}, []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        gid, label, ko, en, tech, _tech_name, note = (list(row) + [None] * 7)[:7]
        if not gid and not label:
            continue
        gid, label = str(gid or "").strip(), str(label or "").strip()
        if not re.fullmatch(r"[a-z0-9]+(-[a-z0-9]+)*", gid):
            errors.append(f"row {r}: group id '{gid}' (lowercase letters, digits, hyphens)")
            continue
        if any(g["id"] == gid for g in groups):
            errors.append(f"row {r}: group id '{gid}' twice")
        # The search words are the two term columns; 대표어 is the group's name
        # (it searches only when it is also listed as a term).
        terms = []
        for term in [*split_terms(ko), *split_terms(en)]:
            if term and key(term) not in {key(t) for t in terms}:
                terms.append(term)
        old = previous.get(gid)
        if old and {key(t) for t in old["terms"]} == {key(t) for t in terms}:
            terms = old["terms"]  # unchanged: keep the stored order and spellings
        if len(terms) < 2:
            errors.append(f"row {r}: '{label}' has fewer than two terms")
        for term in terms:
            if key(term) in owner and owner[key(term)] != gid:
                errors.append(f"row {r}: '{term}' is also in group '{owner[key(term)]}'")
            owner[key(term)] = gid
        group = {"id": gid, "label": label}
        tech = str(tech or "").strip()
        if tech:
            group["techId"] = tech.zfill(2)
        group["terms"] = terms
        if note and str(note).strip():
            group["note"] = str(note).strip()
        groups.append(group)
    rules = json.loads(RULES.read_text(encoding="utf-8"))
    for topic in rules["topics"]:
        for gid in topic["groups"]:
            if gid not in {g["id"] for g in groups}:
                errors.append(f"topic '{topic['label']}' uses group '{gid}', which is gone")
    if errors:
        sys.exit("import refused:\n  " + "\n  ".join(errors))
    added = [g["id"] for g in groups if g["id"] not in previous]
    removed = [gid for gid in previous if gid not in {g["id"] for g in groups}]
    changed = [g["id"] for g in groups if g["id"] in previous and g != previous[g["id"]]]
    print(f"{len(groups)} groups · added {added or '-'} · removed {removed or '-'} · changed {changed or '-'}")
    if check:
        return
    SYNONYMS.write_text(dump_synonyms(syn, groups), encoding="utf-8")
    print(f"wrote {SYNONYMS.relative_to(ROOT)}; next: node scripts/v170/build-search-v170.mjs")


if __name__ == "__main__":
    if len(sys.argv) < 3 or sys.argv[1] not in ("export", "import"):
        sys.exit(__doc__)
    if sys.argv[1] == "export":
        export(Path(sys.argv[2]))
    else:
        import_(Path(sys.argv[2]), "--check" in sys.argv[3:])
