"""Write the V159 contractor handoff workbook from a JSON payload.

Input: a JSON file mapping sheet name -> {"headers": [...], "rows": [[...], ...]}.
Output: an .xlsx with one sheet per key, header row bold + frozen, wrapped text,
and column widths sized to content. No value is computed or rewritten here -
the Node builder (build-typology-handoff-v159.mjs) owns every transformation;
this script only lays cells out.
"""
from __future__ import annotations

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 10


def sheet_width(header: str, rows: list, index: int) -> int:
    lengths = [len(str(header))]
    for row in rows:
        value = row[index] if index < len(row) else ""
        # Widen for the longest *line* of a wrapped cell, not the raw length.
        for line in str(value).split("\n"):
            lengths.append(len(line))
    return max(MIN_COLUMN_WIDTH, min(MAX_COLUMN_WIDTH, max(lengths) + 2))


def main() -> None:
    payload_path, out_path = sys.argv[1], sys.argv[2]
    with open(payload_path, "r", encoding="utf-8") as f:
        sheets = json.load(f)

    wb = Workbook()
    wb.remove(wb.active)
    for name, sheet in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        headers = sheet["headers"]
        rows = sheet["rows"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in rows:
            ws.append(row)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for index, header in enumerate(headers):
            ws.column_dimensions[get_column_letter(index + 1)].width = sheet_width(header, rows, index)
        ws.freeze_panes = "A2"

    wb.save(out_path)
    print(json.dumps({"sheets": {name: len(sheet["rows"]) for name, sheet in sheets.items()}}))


if __name__ == "__main__":
    main()
