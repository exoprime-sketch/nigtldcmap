#!/usr/bin/env python
"""Extract map-display evidence from delivered source workbooks for a country.

Reads the standardized delivery workbooks (sheets 1.1_observation(측정값) /
1.2_entity(레코드형) / 2_meta_info / db_framework) and produces a single
deterministic JSON evidence file that `build_map_candidates_v158.cjs` uses to
judge every framework element against the three map-selection criteria
(see reports/v158/bgd-map-candidates-v158.md for the criteria text).

This script never assumes a fixed country. The country ISO3 comes from
--country; the display name and ADM1 count/label come from
public/data/countries.json; the raw workbook directory comes from
tools/etl/countries/<iso3>/country.json (`source.directory`, resolved
relative to the repository root); the framework size and element labels come
from the workbooks' own db_framework sheet; the "is this coordinate inside
the country" check uses the country's own dissolved outline geometry (found
via its geometry-manifest.json, kind "country-outline"), buffered by
OUTLINE_BUFFER_DEG to absorb coastline generalisation - not a bounding box,
which for a delta country like Bangladesh would also cover neighbouring
India and open sea.

Security: at least one delivered workbook (E-008) carries a live API key in
its `api_params` column. Column lookups are built from each sheet's own
header row and any column whose name contains "api", "key", "token",
"secret", or "password" (case-insensitive) is dropped from the lookup table
before any value is ever read - so a value from such a column can never reach
this script's output, an exception message, or stdout, regardless of which
position it happens to occupy in a given workbook.

Usage:
    python -B tools/etl/countries/bgd/extract_map_evidence.py --country bgd
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from shapely.geometry import Point, shape
from shapely.ops import unary_union
from shapely.prepared import prep

REPO_ROOT = Path(__file__).resolve().parents[4]

# A point-in-country check needs the real coastline, not a bounding box - a
# rectangle over a river-delta country like Bangladesh includes large areas
# that are actually India or open sea. The outline is buffered by this many
# degrees (~5-6 km here) to absorb coastline simplification/generalisation
# before a coordinate is called "outside".
OUTLINE_BUFFER_DEG = 0.05

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

OBSERVATION_SHEET = "1.1_observation(측정값)"
ENTITY_SHEET = "1.2_entity(레코드형)"
META_SHEET = "2_meta_info"
FRAMEWORK_SHEET = "db_framework"

# Never build a column index for a header whose name contains any of these
# substrings (case-insensitive). This is a name-based blocklist, independent
# of cell content, per the handling rule for this delivery.
SECRET_NAME_MARKERS = ("api", "key", "token", "secret", "password")

ELEMENT_ID_RE = re.compile(r"(?<![A-Z0-9])([A-E]-\d{3})(?!\d)")
FILENAME_ELEMENT_ID_RE = re.compile(r"^([A-E]-\d{3})_")

_PLACEHOLDER_EXACT = {
    "", "-", "--", "n/a", "na", "null", "none", "tbd",
    "해당없음", "해당 없음", "미수집", "미제공", "미확보", "자료없음", "자료 없음",
}

# Attribute-label keywords that mark a custom attr_N column as carrying a
# sub-national location/scope (Korean administrative vocabulary varies by
# element, so this is a broad net; the evidence records the matched label
# text so a reviewer can judge each hit).
REGION_ATTR_KEYWORDS = (
    "지역", "행정", "권역", "관구", "유역", "관측소", "station", "basin",
    "division", "district", "province", "region", "p-code", "pcode", "행정코드",
    "주소", "address",
)
# Attribute-label keywords that mark a column as carrying plan/regulation/
# support content (criterion ③'s second question).
POLICY_ATTR_KEYWORDS = (
    "목표", "규제", "지원", "인센티브", "제도", "조항", "보조금", "가격", "한도",
    "세금", "비율", "계획", "발주",
)
NATIONWIDE_TOKENS = {"국가", "전국", "nation", "national", "bgd", "전국/bgd", "national/bgd"}

# Free-text terms worth surfacing from meta_info's rights/notes columns. Purely
# a search list over already-safe columns (license_code, redistribution_allowed,
# download_allowed, 유의사항, source_org) - never over api/key/token/secret ones.
NOTABLE_TERMS = (
    "ODbL", "share-alike", "동일조건변경허락", "표출 불가", "다운로드.*불가",
    "재배포", "GADM", "Mymensingh", "미수록", "보류", "확인필요", "BNEF",
    "구독형", "paywall", "페이월", "India", "인도", "시작점", "중간 정점",
    "대표점", "대표 좌표", "단일 지점", "사용 불가", "저작권", "라이선스",
    "우선순위", "품질", "오차",
)

MAX_ROWS = 250_000
MAX_COLS = 300


# ---------------------------------------------------------------------------
# Small generic helpers (self-contained - no dependency on other ETL modules,
# so this tool carries its own safety guarantees end to end).
# ---------------------------------------------------------------------------


def nfc(value: Any) -> str:
    import unicodedata

    text = str(value)
    return unicodedata.normalize("NFC", text).strip()


def is_placeholder(value: Any) -> bool:
    if value is None:
        return True
    return nfc(value).lower() in _PLACEHOLDER_EXACT


def to_float(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = nfc(value)
    try:
        return float(text)
    except ValueError:
        return None


def extract_element_id(value: Any) -> str | None:
    if value is None:
        return None
    match = ELEMENT_ID_RE.search(nfc(value).upper())
    return match.group(1) if match else None


def is_secret_header(name: str) -> bool:
    lowered = name.lower()
    return any(marker in lowered for marker in SECRET_NAME_MARKERS)


def trim_row(row: Iterable[Any]) -> list[Any]:
    values = list(row)
    while values and values[-1] is None:
        values.pop()
    return values


def read_all_rows(ws: Any) -> list[list[Any]]:
    max_row = min(int(ws.max_row or 0), MAX_ROWS)
    max_col = min(int(ws.max_column or 0), MAX_COLS)
    if max_row <= 0 or max_col <= 0:
        return []
    return [
        trim_row(row)
        for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True)
    ]


def find_header_row(rows: list[list[Any]], marker: str, limit: int = 15) -> int | None:
    for idx, row in enumerate(rows[:limit]):
        if row and row[0] is not None and nfc(row[0]) == marker:
            return idx
    return None


def build_column_map(header_row: list[Any]) -> tuple[dict[str, int], list[str]]:
    """Map header name -> column index, dropping any secret-named column.

    A dropped column has no entry, so no later code path can read its value
    by name - this is enforced structurally, not by remembering to check.
    """

    col_map: dict[str, int] = {}
    redacted: list[str] = []
    for idx, raw_name in enumerate(header_row):
        if raw_name is None:
            continue
        name = nfc(raw_name)
        if not name:
            continue
        if is_secret_header(name):
            redacted.append(name)
            continue
        col_map.setdefault(name, idx)
    return col_map, redacted


def cell(row: list[Any], col_map: dict[str, int], name: str) -> Any:
    idx = col_map.get(name)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def is_data_row(row: list[Any], col_map: dict[str, int]) -> bool:
    """A real data row has an extractable element_id and a populated indicator_id.

    This also naturally excludes the sheet's own Korean label sub-header row
    (whose first cell repeats the literal string "요소_id") and single-cell
    explanatory placeholder rows, without assuming a fixed row offset.
    """

    eid = extract_element_id(cell(row, col_map, "요소_id"))
    if not eid:
        return False
    if "indicator_id" not in col_map:
        return True
    return not is_placeholder(cell(row, col_map, "indicator_id"))


def read_sheet(wb: Any, sheet_name: str, marker: str = "요소_id") -> dict[str, Any]:
    if sheet_name not in wb.sheetnames:
        return {"present": False, "colMap": {}, "redactedColumns": [], "labelRow": [], "dataRows": []}
    rows = read_all_rows(wb[sheet_name])
    header_idx = find_header_row(rows, marker)
    if header_idx is None:
        return {"present": True, "colMap": {}, "redactedColumns": [], "labelRow": [], "dataRows": [], "headerNotFound": True}
    header_row = rows[header_idx]
    col_map, redacted = build_column_map(header_row)
    label_row = rows[header_idx + 1] if header_idx + 1 < len(rows) else []
    data_rows = [row for row in rows[header_idx + 1 :] if is_data_row(row, col_map)]
    return {
        "present": True,
        "colMap": col_map,
        "redactedColumns": redacted,
        "labelRow": label_row if not is_data_row(label_row, col_map) else [],
        "dataRows": data_rows,
    }


def attribute_labels_from(sheet: dict[str, Any]) -> dict[str, str]:
    labels: dict[str, str] = {}
    label_row = sheet["labelRow"]
    for name, idx in sheet["colMap"].items():
        if not name.startswith("attr_"):
            continue
        if idx < len(label_row) and label_row[idx]:
            labels[name] = nfc(label_row[idx])
    return labels


def keyword_hits(text: str, keywords: tuple[str, ...]) -> list[str]:
    lowered = text.lower()
    return [kw for kw in keywords if kw.lower() in lowered]


def normalize_token(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def classify_region_token(value: str, division_keys: set[str], district_keys: set[str]) -> str | None:
    token = normalize_token(value)
    if not token:
        return None
    if token in division_keys:
        return "division"
    if token in district_keys:
        return "district"
    return None


def find_notable_flags(text: str, terms: tuple[str, ...], max_flags: int = 8) -> list[dict[str, str]]:
    flags: list[dict[str, str]] = []
    for term in terms:
        try:
            match = re.search(term, text, re.IGNORECASE)
        except re.error:
            continue
        if not match:
            continue
        start = max(0, match.start() - 50)
        end = min(len(text), match.end() + 90)
        snippet = re.sub(r"\s+", " ", text[start:end]).strip()
        flags.append({"term": term, "snippet": snippet})
        if len(flags) >= max_flags:
            break
    return flags


# ---------------------------------------------------------------------------
# Registry / configuration loading (no hardcoded country name, ISO3 or counts)
# ---------------------------------------------------------------------------


def load_country_registry(iso3: str) -> dict[str, Any]:
    path = REPO_ROOT / "public" / "data" / "countries.json"
    doc = json.loads(path.read_text(encoding="utf-8"))
    for entry in doc["countries"]:
        if entry["iso3"].upper() == iso3.upper():
            return entry
    raise SystemExit(f"country {iso3} not found in {path}")


def load_country_etl_config(iso3: str) -> dict[str, Any]:
    path = REPO_ROOT / "tools" / "etl" / "countries" / iso3.lower() / "country.json"
    return json.loads(path.read_text(encoding="utf-8"))


def load_region_key_sets(iso3: str) -> tuple[set[str], set[str]]:
    """Division- and district-level key sets from the platform's own region
    name dictionary (src/data/geo/regionNamesV161.json) - never hardcoded here.
    """

    path = REPO_ROOT / "src" / "data" / "geo" / "regionNamesV161.json"
    doc = json.loads(path.read_text(encoding="utf-8"))
    block = doc.get("countries", {}).get(iso3.upper(), {})
    division_keys: set[str] = set()
    district_keys: set[str] = set()
    for entry in block.get("entries", []):
        keys = {normalize_token(k) for k in entry.get("keys", [])}
        if entry.get("level") == "division":
            division_keys |= keys
        elif entry.get("level") == "district":
            district_keys |= keys
    return division_keys, district_keys


def load_country_outline_prepared(iso3: str, data_root: str) -> Any:
    """The country's dissolved outline (from its own geometry manifest, never
    a hardcoded filename), buffered by OUTLINE_BUFFER_DEG and wrapped with
    shapely.prepared.prep for fast repeated point-in-polygon tests.
    """

    manifest_path = REPO_ROOT / "public" / data_root.lstrip("/") / "geometry" / "geometry-manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    outline_asset = next((asset for asset in manifest["assets"] if asset.get("kind") == "country-outline"), None)
    if outline_asset is None:
        raise SystemExit(f"no country-outline asset in {manifest_path}")
    outline_path = REPO_ROOT / "public" / outline_asset["url"].lstrip("/")
    geojson = json.loads(outline_path.read_text(encoding="utf-8"))
    geometries = [shape(feature["geometry"]) for feature in geojson["features"]]
    outline = geometries[0] if len(geometries) == 1 else unary_union(geometries)
    return prep(outline.buffer(OUTLINE_BUFFER_DEG))


def list_delivered_workbooks(source_dir: Path) -> dict[str, Path]:
    delivered: dict[str, Path] = {}
    for entry in sorted(source_dir.iterdir()):
        if not entry.is_file() or entry.suffix.lower() != ".xlsx":
            continue
        if entry.name.startswith("._") or entry.name.startswith("~$"):
            continue
        match = FILENAME_ELEMENT_ID_RE.match(entry.name)
        if not match:
            continue
        delivered[match.group(1)] = entry
    return delivered


def read_framework(delivered: dict[str, Path]) -> dict[str, dict[str, Any]]:
    """Read the universal 152-element framework from the first delivered
    workbook (sorted) whose db_framework sheet parses. Every workbook in the
    delivery carries the same framework listing (see CLAUDE.md / country.json
    note); we do not assume which specific code it is.
    """

    for element_id in sorted(delivered):
        path = delivered[element_id]
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if FRAMEWORK_SHEET not in wb.sheetnames:
                continue
            rows = read_all_rows(wb[FRAMEWORK_SHEET])
            header_idx = find_header_row(rows, "요소ID")
            if header_idx is None:
                continue
            col_map, _ = build_column_map(rows[header_idx])
            framework: dict[str, dict[str, Any]] = {}
            for row in rows[header_idx + 1 :]:
                eid = extract_element_id(cell(row, col_map, "요소ID"))
                if not eid:
                    continue
                framework[eid] = {
                    "category": nfc(cell(row, col_map, "대분류_CD") or ""),
                    "categoryLabel": nfc(cell(row, col_map, "대분류_KR") or ""),
                    "subCode": nfc(cell(row, col_map, "중분류_CD") or ""),
                    "subLabel": nfc(cell(row, col_map, "중분류_KR") or ""),
                    "fineCode": nfc(cell(row, col_map, "세분류_CD") or ""),
                    "fineLabel": nfc(cell(row, col_map, "세분류_KR") or ""),
                    "label": nfc(cell(row, col_map, "요소") or ""),
                }
            if len(framework) >= 100:
                return framework
        finally:
            wb.close()
    raise SystemExit("could not read a usable db_framework sheet from any delivered workbook")


# ---------------------------------------------------------------------------
# Per-element evidence
# ---------------------------------------------------------------------------


def analyze_indicator_suffixes(
    element_id: str,
    indicator_ids: Iterable[str],
    division_keys: set[str],
    district_keys: set[str],
) -> dict[str, Any]:
    prefix = f"{element_id}_"
    single_token_matches: dict[str, int] = {}
    base_groups: dict[str, set[str]] = {}
    composite_hits = 0
    total = 0
    for indicator_id in indicator_ids:
        if not indicator_id or not indicator_id.startswith(prefix):
            continue
        total += 1
        tokens = indicator_id[len(prefix) :].split("_")
        if not tokens:
            continue
        last_norm = normalize_token(tokens[-1])
        if last_norm in division_keys or last_norm in district_keys:
            single_token_matches[last_norm] = single_token_matches.get(last_norm, 0) + 1
            base = "_".join(tokens[:-1])
            base_groups.setdefault(base, set()).add(last_norm)
        for window in (2, 3, 4):
            if len(tokens) >= window:
                trailing = [normalize_token(t) for t in tokens[-window:]]
                if all(t in district_keys for t in trailing):
                    composite_hits += 1
                    break
    distinct_tokens = set(single_token_matches)
    groups_multi = {base: sorted(toks) for base, toks in base_groups.items() if len(toks) >= 2}
    division_only_tokens = sorted(distinct_tokens & division_keys)
    district_only_tokens = sorted(distinct_tokens & district_keys)
    return {
        "distinctMatchedTokens": sorted(distinct_tokens),
        "divisionTokensMatched": division_only_tokens,
        "districtTokensMatched": district_only_tokens,
        "allMatchedTokensAreDivisionNames": bool(distinct_tokens) and distinct_tokens.issubset(division_keys),
        # base indicator name -> sorted region tokens it varies across (size shows
        # whether the grouping looks like our 8-division scheme or something else,
        # e.g. Global Data Lab's larger, differently-shaped region groups).
        "indicatorBaseRegionVariety": groups_multi,
        "maxDistinctTokensPerBase": max((len(toks) for toks in groups_multi.values()), default=0),
        "compositeMultiDistrictSuffixCount": composite_hits,
        "compositeMultiDistrictSuffixRatio": round(composite_hits / total, 4) if total else 0.0,
        "totalIndicatorsChecked": total,
    }


def analyze_element(
    element_id: str,
    path: Path,
    outline_prepared: Any,
    division_keys: set[str],
    district_keys: set[str],
) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        warnings: list[str] = []
        obs = read_sheet(wb, OBSERVATION_SHEET)
        ent = read_sheet(wb, ENTITY_SHEET)
        meta = read_sheet(wb, META_SHEET)
        for label, sheet in (("observation", obs), ("entity", ent), ("meta", meta)):
            if not sheet["present"]:
                warnings.append(f"missing-sheet:{label}")
            elif sheet.get("headerNotFound"):
                warnings.append(f"header-not-found:{label}")

        redacted_columns = sorted(set(obs["redactedColumns"] + ent["redactedColumns"] + meta["redactedColumns"]))

        obs_indicator_ids = sorted({nfc(cell(r, obs["colMap"], "indicator_id")) for r in obs["dataRows"] if cell(r, obs["colMap"], "indicator_id")})
        ent_indicator_ids = sorted({nfc(cell(r, ent["colMap"], "indicator_id")) for r in ent["dataRows"] if cell(r, ent["colMap"], "indicator_id")})

        # --- coordinates & geometry (entity sheet) ---
        # "Outside" is judged against the real country outline (buffered to
        # absorb coastline generalisation), not a bounding box - a box over a
        # delta country like Bangladesh also covers neighbouring India and
        # open sea, so it cannot tell a real border point from an error.
        coord_rows = 0
        outside_outline_count = 0
        outside_outline_samples: list[dict[str, Any]] = []
        geometry_types: set[str] = set()
        for row in ent["dataRows"]:
            lat = to_float(cell(row, ent["colMap"], "lat"))
            lon = to_float(cell(row, ent["colMap"], "lon"))
            gtype = cell(row, ent["colMap"], "geometry_type")
            if gtype and not is_placeholder(gtype):
                geometry_types.add(nfc(gtype))
            if lat is None or lon is None:
                continue
            coord_rows += 1
            if not outline_prepared.contains(Point(lon, lat)):
                outside_outline_count += 1
                if len(outside_outline_samples) < 40:
                    outside_outline_samples.append(
                        {
                            "indicatorId": nfc(cell(row, ent["colMap"], "indicator_id") or ""),
                            "lat": lat,
                            "lon": lon,
                            "attr_1": nfc(cell(row, ent["colMap"], "attr_1") or "") or None,
                            "attr_2": nfc(cell(row, ent["colMap"], "attr_2") or "") or None,
                            "attr_3": nfc(cell(row, ent["colMap"], "attr_3") or "") or None,
                            "attr_4": nfc(cell(row, ent["colMap"], "attr_4") or "") or None,
                        }
                    )

        # --- descriptive / region-like attribute columns (entity sheet) ---
        attribute_labels = attribute_labels_from(ent)
        region_like_attrs: dict[str, Any] = {}
        declared_region_values: set[str] = set()
        for name, label in attribute_labels.items():
            hits = keyword_hits(label, REGION_ATTR_KEYWORDS)
            if not hits:
                continue
            idx = ent["colMap"].get(name)
            values: set[str] = set()
            if idx is not None:
                for row in ent["dataRows"]:
                    value = row[idx] if idx < len(row) else None
                    if value is not None and not is_placeholder(value):
                        values.add(nfc(value))
            declared_region_values |= values
            division_hits = sorted({v for v in values if classify_region_token(v, division_keys, district_keys) == "division"})
            district_hits = sorted({v for v in values if classify_region_token(v, division_keys, district_keys) == "district"})
            region_like_attrs[name] = {
                "label": label,
                "matchedKeywords": hits,
                "distinctValueCount": len(values),
                "distinctValuesSample": sorted(values)[:80],
                "divisionValuesMatched": division_hits,
                "districtValuesMatched": district_hits,
            }
        policy_like_attrs = {
            name: {"label": label, "matchedKeywords": keyword_hits(label, POLICY_ATTR_KEYWORDS)}
            for name, label in attribute_labels.items()
            if keyword_hits(label, POLICY_ATTR_KEYWORDS)
        }
        declared_region_classified = {
            "division": sorted({v for v in declared_region_values if classify_region_token(v, division_keys, district_keys) == "division"}),
            "district": sorted({v for v in declared_region_values if classify_region_token(v, division_keys, district_keys) == "district"}),
            "unclassified": sorted(
                v for v in declared_region_values if classify_region_token(v, division_keys, district_keys) is None and v.lower() not in NATIONWIDE_TOKENS
            )[:60],
        }
        non_nationwide_declared = sorted(v for v in declared_region_values if v.lower() not in NATIONWIDE_TOKENS)

        populated_attr_columns = sorted(
            {
                name
                for name, idx in ent["colMap"].items()
                if name.startswith("attr_") and any(idx < len(row) and not is_placeholder(row[idx]) for row in ent["dataRows"])
            }
        )
        has_note_content = any(not is_placeholder(cell(row, ent["colMap"], "note")) for row in ent["dataRows"])

        # --- indicator suffix analysis (division/district comparison) ---
        obs_suffixes = analyze_indicator_suffixes(element_id, obs_indicator_ids, division_keys, district_keys)
        ent_suffixes = analyze_indicator_suffixes(element_id, ent_indicator_ids, division_keys, district_keys)

        # --- meta_info: spatial units + rights/notes ---
        spatial_units: set[str] = set()
        source_orgs: set[str] = set()
        license_notes: dict[tuple, dict[str, Any]] = {}
        for row in meta["dataRows"]:
            su = cell(row, meta["colMap"], "spatial_unit")
            if su and not is_placeholder(su):
                spatial_units.add(nfc(su))
            org = cell(row, meta["colMap"], "source_org")
            if org and not is_placeholder(org):
                source_orgs.add(nfc(org))
            indicator_id = nfc(cell(row, meta["colMap"], "indicator_id") or "")
            license_code = nfc(cell(row, meta["colMap"], "license_code") or "") or None
            redistribution = nfc(cell(row, meta["colMap"], "redistribution_allowed") or "") or None
            download = nfc(cell(row, meta["colMap"], "download_allowed") or "") or None
            note = nfc(cell(row, meta["colMap"], "유의사항") or "")
            key = (license_code, redistribution, download, note)
            if key not in license_notes:
                license_notes[key] = {
                    "indicatorIds": [],
                    "spatialUnit": nfc(su or "") or None,
                    "licenseCode": license_code,
                    "redistributionAllowed": redistribution,
                    "downloadAllowed": download,
                    "sourceOrg": nfc(org or "") or None,
                    "notePrefix": note[:280] if note else None,
                    "flags": find_notable_flags(note, NOTABLE_TERMS) if note else [],
                }
            if indicator_id and indicator_id not in license_notes[key]["indicatorIds"]:
                license_notes[key]["indicatorIds"].append(indicator_id)

        division_keys_found = sorted(
            set(declared_region_classified["division"])
            | set(obs_suffixes["divisionTokensMatched"])
            | set(ent_suffixes["divisionTokensMatched"])
        )
        district_keys_found = sorted(
            set(declared_region_classified["district"])
            | set(obs_suffixes["districtTokensMatched"])
            | set(ent_suffixes["districtTokensMatched"])
        )
        has_subnational_variation = (
            coord_rows > 0
            or bool(non_nationwide_declared)
            or bool(obs_suffixes["indicatorBaseRegionVariety"])
            or bool(ent_suffixes["indicatorBaseRegionVariety"])
        )
        nationwide_only = bool(spatial_units) and spatial_units.issubset({"nation"}) and not has_subnational_variation

        return {
            "archiveName": path.name,
            "warnings": warnings,
            "redactedColumns": redacted_columns,
            "sheets": {
                "observation": {"dataRowCount": len(obs["dataRows"]), "uniqueIndicatorCount": len(obs_indicator_ids), "indicatorIdSample": obs_indicator_ids[:40]},
                "entity": {"dataRowCount": len(ent["dataRows"]), "uniqueIndicatorCount": len(ent_indicator_ids), "indicatorIdSample": ent_indicator_ids[:40]},
                "meta": {"dataRowCount": len(meta["dataRows"])},
            },
            "spatialUnits": sorted(spatial_units),
            "sourceOrganizations": sorted(source_orgs),
            "nationwideOnly": nationwide_only,
            "coordinates": {
                "rowsWithLatLon": coord_rows,
                "outsideOutlineCount": outside_outline_count if coord_rows else 0,
                "outsideOutlineBufferDeg": OUTLINE_BUFFER_DEG,
                "outsideOutlineSamples": outside_outline_samples,
            },
            "geometryTypes": sorted(geometry_types),
            "regionEvidence": {
                "regionLikeAttrs": region_like_attrs,
                "policyLikeAttrs": policy_like_attrs,
                "declaredRegionValueCount": len(declared_region_values),
                "declaredRegionValuesNonNationwideSample": non_nationwide_declared[:60],
                "declaredRegionClassified": declared_region_classified,
                "observationIndicatorSuffixes": obs_suffixes,
                "entityIndicatorSuffixes": ent_suffixes,
                # Rollup used by the candidate judgment: the union of every
                # division/district match this element produced, from either the
                # region-like attribute columns or the indicator_id suffixes.
                "divisionKeysFound": division_keys_found,
                "divisionKeysFoundCount": len(division_keys_found),
                "districtKeysFound": district_keys_found,
                "districtKeysFoundCount": len(district_keys_found),
                "hasSubnationalVariation": has_subnational_variation,
            },
            "descriptiveAttributes": {
                "populatedAttrColumns": populated_attr_columns,
                "attributeLabels": attribute_labels,
                "hasNoteContent": has_note_content,
            },
            "licenseNotes": list(license_notes.values())[:20],
        }
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--country", required=True, help="ISO3 country code, e.g. bgd")
    parser.add_argument("--out", default=None, help="output JSON path (default: reports/v158/<iso3>-map-evidence-v158.json)")
    parser.add_argument("--only", default=None, help="comma-separated element IDs to limit processing to (debugging)")
    args = parser.parse_args()

    iso3 = args.country.upper()
    registry_entry = load_country_registry(iso3)
    etl_config = load_country_etl_config(iso3)
    division_keys, district_keys = load_region_key_sets(iso3)
    outline_prepared = load_country_outline_prepared(iso3, registry_entry["dataRoot"])

    source_dir = (REPO_ROOT / etl_config["source"]["directory"]).resolve()
    if not source_dir.is_dir():
        raise SystemExit(f"source directory not found: {source_dir}")

    delivered = list_delivered_workbooks(source_dir)
    if args.only:
        wanted = {code.strip().upper() for code in args.only.split(",") if code.strip()}
        delivered = {k: v for k, v in delivered.items() if k in wanted}

    framework = read_framework(list_delivered_workbooks(source_dir))
    not_provided = sorted(set(framework) - set(list_delivered_workbooks(source_dir)))

    category_counts: dict[str, int] = {}
    for info in framework.values():
        category_counts[info["category"]] = category_counts.get(info["category"], 0) + 1
    delivered_category_counts: dict[str, int] = {}
    for element_id in list_delivered_workbooks(source_dir):
        cat = framework.get(element_id, {}).get("category", element_id[0])
        delivered_category_counts[cat] = delivered_category_counts.get(cat, 0) + 1

    elements: dict[str, Any] = {}
    for element_id in sorted(delivered):
        path = delivered[element_id]
        print(f"  analyzing {element_id} ({path.name})", file=sys.stderr)
        evidence = analyze_element(element_id, path, outline_prepared, division_keys, district_keys)
        evidence["framework"] = framework.get(element_id, {})
        elements[element_id] = evidence

    output = {
        "schemaVersion": "bgd-map-evidence-v158",
        "generatedFor": iso3,
        "country": {
            "iso3": iso3,
            "nameKo": registry_entry["nameKo"],
            "nameEn": registry_entry["nameEn"],
            "adm1Count": registry_entry["adm"]["level1"]["count"],
            "adm1Label": registry_entry["adm"]["level1"]["label"],
            "bbox": registry_entry["bbox"],
        },
        "sourceDirectory": etl_config["source"]["directory"],
        "framework": {
            "totalElements": len(framework),
            "deliveredCount": len(list_delivered_workbooks(source_dir)),
            "notProvidedCount": len(not_provided),
            "notProvidedCodes": not_provided,
            "categoryCounts": dict(sorted(category_counts.items())),
            "deliveredCategoryCounts": dict(sorted(delivered_category_counts.items())),
            # Every one of the 152 codes (delivered or not) with its label, so a
            # reviewer can look up the "대상 데이터" text for not-provided codes too.
            "allElements": framework,
        },
        "regionDictionary": {
            "source": "src/data/geo/regionNamesV161.json",
            "divisionKeyCount": len(division_keys),
            "districtKeyCount": len(district_keys),
        },
        "elements": elements,
    }

    out_path = Path(args.out) if args.out else REPO_ROOT / "reports" / "v158" / f"{iso3.lower()}-map-evidence-v158.json"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(output, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(f"wrote {out_path} ({len(elements)} delivered elements, {len(not_provided)} not provided)", file=sys.stderr)


if __name__ == "__main__":
    main()
